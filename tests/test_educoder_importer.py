from openpyxl import Workbook

from models import StudentPractice
from services.importers.educoder_importer import EducoderImporter


def test_total_score_is_personal_total_averaged_over_practice_training_columns(tmp_path):
    """实践成绩应按个人总成绩在实训子列间折算为百分制。"""
    file_path = tmp_path / '头歌-总成绩.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '学生总成绩'

    sheet.append(['姓名', '学号', '个人总成绩', '实践训练', None, None, '在线考试'])
    sheet.append([None, None, None, '实训一', '实训二', '实训三', None])
    sheet.append(['测试学生', '20260001', 240, 100, 80, 60, 90])
    sheet.merge_cells('A1:A2')
    sheet.merge_cells('B1:B2')
    sheet.merge_cells('C1:C2')
    sheet.merge_cells('D1:F1')
    sheet.merge_cells('G1:G2')
    workbook.save(file_path)

    importer = EducoderImporter(str(file_path), 1, 'admin')
    importer.df = importer._read_file()

    parsed, errors = importer.parse()

    assert parsed is True
    assert errors == []
    assert importer.parsed_data == [{
        'student_no': '20260001',
        'name': '测试学生',
        'total_score': 80.0,
        'personal_total_score': 240.0,
        'practice_training_count': 3,
    }]


def test_practice_score_uses_the_normalized_educoder_total_score():
    """页面展示的实践评分应与头歌总成绩折算结果一致。"""
    practice = StudentPractice(
        total_score=80,
        activity_score=100,
        avg_experiment_score=50,
    )

    assert practice.calculate_practice_score() == 80.0
